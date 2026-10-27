#!/usr/bin/env python3
"""Wandelt die Markdown-Vorlagen in druckfertige Word-Dateien um (Vorlagen/Word/).

Aufruf:  python3 Bauleitung/_werkzeuge/vorlagen_zu_word.py
Braucht: pandoc, python-docx
"""
import pathlib
import subprocess

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = pathlib.Path(__file__).resolve().parent.parent
ZIEL = ROOT / "Vorlagen" / "Word"
QUELLEN = (
    sorted((ROOT / "Vorlagen").glob("*.md"))
    + sorted((ROOT / "00_Notfall").glob("*.md"))
    + [ROOT / "04_Baustellenverordnung" / "Vorankuendigung_Anhang_I.md", ROOT / "Checkliste" / "CHECKLISTE.md"]
)
QUER_AB_SPALTEN = 7  # breite Tabellen → Querformat


def rahmen(tabelle):
    tbl_pr = tabelle._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for kante in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{kante}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), "555555")
        borders.append(el)
    tbl_pr.append(borders)


def kopf_schattieren(zeile):
    for zelle in zeile.cells:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "E8E8E8")
        zelle._tc.get_or_add_tcPr().append(shd)
        for absatz in zelle.paragraphs:
            for lauf in absatz.runs:
                lauf.bold = True


def mindesthoehe(zeile, cm):
    tr_pr = zeile._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(cm * 567)))
    h.set(qn("w:hRule"), "atLeast")
    tr_pr.append(h)


def nachbearbeiten(pfad):
    doc = Document(pfad)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    for name in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        if name in doc.styles:
            s = doc.styles[name]
            s.font.name = "Arial"
            s.font.color.rgb = RGBColor(0x23, 0x28, 0x2C)
    if "Compact" in doc.styles:
        doc.styles["Compact"].font.size = Pt(10)

    quer = any(len(t.columns) >= QUER_AB_SPALTEN for t in doc.tables)
    for sec in doc.sections:
        if quer:
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
        else:
            sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        for rand in ("left_margin", "right_margin"):
            setattr(sec, rand, Cm(1.8))
        sec.top_margin = sec.bottom_margin = Cm(1.5)

    for t in doc.tables:
        rahmen(t)
        if t.rows:
            kopf_schattieren(t.rows[0])
        for zeile in t.rows[1:]:
            mindesthoehe(zeile, 0.75)
    doc.save(pfad)


def main():
    ZIEL.mkdir(parents=True, exist_ok=True)
    for quelle in QUELLEN:
        ziel = ZIEL / (quelle.stem + ".docx")
        subprocess.run(
            ["pandoc", str(quelle), "-f", "markdown+pipe_tables+task_lists", "-t", "docx", "-o", str(ziel)],
            check=True,
        )
        nachbearbeiten(ziel)
        print("✓", ziel.relative_to(ROOT))


if __name__ == "__main__":
    main()
